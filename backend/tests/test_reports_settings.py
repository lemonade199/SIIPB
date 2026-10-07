"""Laporan PDF/Excel, dashboard, pengaturan (SMTP terenkripsi), audit log, backup."""
import io

import pytest


@pytest.mark.parametrize("rtype", ["inventaris", "peminjaman", "pengembalian", "keterlambatan", "kerusakan"])
def test_reports_all_formats(client, auth_headers, make_borrowing, rtype):
    make_borrowing()
    js = client.get(f"/api/v1/reports/{rtype}", headers=auth_headers)
    assert js.status_code == 200 and js.get_json()["data"]["columns"]
    x = client.get(f"/api/v1/reports/{rtype}?format=xlsx", headers=auth_headers)
    assert x.status_code == 200 and x.data[:2] == b"PK"
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(x.data))
    assert wb.active.cell(row=5, column=1).value == js.get_json()["data"]["columns"][0]["label"]
    p = client.get(f"/api/v1/reports/{rtype}?format=pdf", headers=auth_headers)
    assert p.status_code == 200 and p.data.startswith(b"%PDF")


def test_report_filters_and_permission(client, auth_headers, pimpinan_headers):
    rows = client.get("/api/v1/reports/inventaris?status=DIPINJAM", headers=auth_headers).get_json()["data"]["rows"]
    assert all(r["status"] == "DIPINJAM" for r in rows)
    assert client.get("/api/v1/reports/inventaris?format=pdf", headers=pimpinan_headers).status_code == 200
    assert client.get("/api/v1/reports/tidak-ada", headers=auth_headers).status_code == 404


def test_dashboard(client, auth_headers):
    s = client.get("/api/v1/dashboard/summary", headers=auth_headers).get_json()["data"]
    assert {"assets", "borrowings", "incidents"} <= set(s)
    st = client.get("/api/v1/dashboard/statistics?months=6", headers=auth_headers).get_json()["data"]
    assert len(st["monthly"]) == 6 and "top_items" in st


def test_settings_smtp_password_encrypted(client, auth_headers, petugas_headers, db_session):
    res = client.put("/api/v1/settings", json={"institution": "Instansi Uji", "smtp": {"host": "smtp.contoh.id", "port": 465,
                                                  "encryption": "SSL/TLS", "password": "sandi-smtp-rahasia"}}, headers=auth_headers)
    assert res.status_code == 200
    data = res.get_json()["data"]
    assert data["institution"] == "Instansi Uji" and data["smtp"]["password_set"] is True
    assert "sandi-smtp-rahasia" not in res.get_data(as_text=True)
    from sqlalchemy import select

    from app.models import SystemSetting
    from app.services.settings_service import decrypt_secret

    row = db_session.scalars(select(SystemSetting).where(SystemSetting.key == "smtp.password")).first()
    assert row.is_secret and row.value != "sandi-smtp-rahasia" and decrypt_secret(row.value) == "sandi-smtp-rahasia"
    bad = client.put("/api/v1/settings", json={"smtp": {"port": 99999}, "scheduler": {"time": "25:00"}}, headers=auth_headers)
    assert bad.status_code == 400 and {"smtp.port", "scheduler.time"} <= set(bad.get_json()["errors"])
    public = client.get("/api/v1/settings", headers=petugas_headers).get_json()["data"]
    assert "smtp" not in public and public["institution"] == "Instansi Uji"
    client.put("/api/v1/settings", json={"smtp": {"host": "smtp.gmail.com", "port": 587, "encryption": "STARTTLS"}}, headers=auth_headers)
    # bersihkan kata sandi agar pengiriman tetap disimulasikan
    db_session.delete(row)
    db_session.commit()


def test_audit_log_records_changes(client, auth_headers, make_asset):
    a = make_asset()
    client.put(f"/api/v1/assets/{a['id']}", json={"name": "Nama Baru"}, headers=auth_headers)
    logs = client.get(f"/api/v1/audit-logs?entity_type=asset&entity_id={a['id']}", headers=auth_headers).get_json()["data"]
    upd = next(lg for lg in logs if lg["action"] == "UPDATE")
    assert upd["old_data"]["name"] == "Barang Uji" and upd["new_data"]["name"] == "Nama Baru"
    assert upd["user_name"] == "Administrator SIIPB"


def test_backup_endpoint(client, auth_headers):
    import shutil

    res = client.post("/api/v1/backups", headers=auth_headers)
    if not (shutil.which("mariadb-dump") or shutil.which("mysqldump")):
        assert res.status_code == 400
        return
    assert res.status_code == 200, res.get_json()
    name = res.get_json()["data"]["file"]
    assert name in [b["file"] for b in client.get("/api/v1/backups", headers=auth_headers).get_json()["data"]]
    dl = client.get(f"/api/v1/backups/{name}", headers=auth_headers)
    assert dl.status_code == 200 and dl.data[:2] == b"\x1f\x8b"
    assert client.get("/api/v1/backups/..%2f..%2fetc%2fpasswd", headers=auth_headers).status_code == 404
