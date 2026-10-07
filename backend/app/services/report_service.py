"""Laporan (PB-013): inventaris, peminjaman, pengembalian, keterlambatan, kerusakan & kehilangan.

Format: JSON (pratinjau), Excel (openpyxl), PDF (WeasyPrint). Kolom sama dengan pratinjau di frontend.
"""
from __future__ import annotations

import html
import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Asset,
    AssetHistory,
    Borrowing,
    BorrowingStatus,
    NotificationEvent,
    Return,
)
from app.services.notification_service import EVENT_TO_UI, fmt_tanggal

REPORT_TYPES = {
    "inventaris": "Laporan Inventaris Barang",
    "peminjaman": "Laporan Peminjaman",
    "pengembalian": "Laporan Pengembalian",
    "keterlambatan": "Laporan Keterlambatan",
    "kerusakan": "Laporan Kerusakan & Kehilangan",
}
CONDITION_LABEL = {"BAIK": "Baik", "RUSAK_RINGAN": "Rusak ringan", "RUSAK_BERAT": "Rusak berat"}
BORROW_LABEL = {"AKTIF": "DIPINJAM", "TERLAMBAT": "TERLAMBAT", "DIKEMBALIKAN": "DIKEMBALIKAN", "DIBATALKAN": "DIBATALKAN", "DRAF": "DRAF"}


@dataclass
class Report:
    type: str
    title: str
    columns: list[dict[str, Any]]
    rows: list[dict[str, Any]]
    summary: str
    filters: dict[str, Any]


def _d(v: date | datetime | None) -> str:
    return fmt_tanggal(v, with_day=False) if v else "—"


def _rupiah(v: float | int | None) -> str:
    return "Rp " + f"{int(round(v or 0)):,}".replace(",", ".")


def _date(v: str | None) -> date | None:
    try:
        return date.fromisoformat(v) if v else None
    except ValueError:
        return None


def build_report(session: Session, rtype: str, f: dict[str, Any], today: date | None = None) -> Report:
    if rtype not in REPORT_TYPES:
        raise ValueError(f"Jenis laporan tidak dikenal: {rtype}")
    today = today or date.today()
    d_from, d_to = _date(f.get("from")), _date(f.get("to"))
    unit_id = int(f["unit_id"]) if f.get("unit_id") else None
    category_id = int(f["category_id"]) if f.get("category_id") else None
    location_id = int(f["location_id"]) if f.get("location_id") else None
    status = f.get("status") or None

    def in_range(d: date) -> bool:
        return (not d_from or d >= d_from) and (not d_to or d <= d_to)

    def asset_ok(a: Asset) -> bool:
        return (not category_id or a.category_id == category_id) and (not location_id or a.location_id == location_id)

    def unit_ok(b: Borrowing) -> bool:
        return not unit_id or b.borrower.unit_id == unit_id

    def items_txt(b: Borrowing) -> str:
        return "; ".join(f"{it.asset.name} ({it.asset.inventory_code})" for it in b.items)

    def counted(b: Borrowing) -> bool:
        return b.status not in (BorrowingStatus.DRAF.value, BorrowingStatus.DIBATALKAN.value)

    def officer(b: Borrowing) -> str:
        u = b.checkout_user or b.handler
        return u.full_name if u else "—"

    def late_days(b: Borrowing) -> int:
        end = max((r.returned_at.date() for r in b.returns), default=None) if b.status == BorrowingStatus.DIKEMBALIKAN.value else today
        return max(0, (end - b.due_date).days) if end else 0

    def display(b: Borrowing) -> str:
        if b.status in (BorrowingStatus.AKTIF.value, BorrowingStatus.TERLAMBAT.value):
            if b.due_date < today:
                return "TERLAMBAT"
            if b.due_date == today:
                return "JATUH TEMPO"
        return BORROW_LABEL.get(b.status, b.status)

    if rtype == "inventaris":
        assets = session.scalars(select(Asset).where(Asset.deleted_at.is_(None), Asset.is_active.is_(True)).order_by(Asset.inventory_code)).all()
        rows = [
            {
                "kode": a.inventory_code,
                "nama": a.name,
                "kategori": a.category.name if a.category else "",
                "merek": " / ".join(x for x in (a.brand, a.model) if x),
                "seri": a.serial_number or "",
                "tahun": a.purchase_date.year if a.purchase_date else "",
                "sumber": a.acquisition_source or "",
                "nilai": float(a.acquisition_cost) if a.acquisition_cost is not None else None,
                "lokasi": a.location.name if a.location else "",
                "kondisi": CONDITION_LABEL.get(a.condition, a.condition),
                "status": a.status,
            }
            for a in assets
            if asset_ok(a) and (not status or a.status == status)
        ]
        cols = [("kode", "Kode"), ("nama", "Nama Barang"), ("kategori", "Kategori"), ("merek", "Merek/Model"),
                ("seri", "No. Seri"), ("tahun", "Tahun"), ("sumber", "Sumber"), ("nilai", "Nilai Perolehan", True),
                ("lokasi", "Lokasi"), ("kondisi", "Kondisi"), ("status", "Status")]
        summary = f"{len(rows)} barang · total nilai {_rupiah(sum(r['nilai'] or 0 for r in rows))}"
    elif rtype == "peminjaman":
        bs = session.scalars(select(Borrowing).order_by(Borrowing.start_date.desc(), Borrowing.id.desc())).all()
        rows = [
            {
                "kode": b.transaction_number,
                "peminjam": b.borrower.name,
                "unit": b.borrower.unit.name if b.borrower.unit else "",
                "barang": items_txt(b),
                "pinjam": _d(b.start_date),
                "batas": _d(b.due_date),
                "status": display(b),
                "petugas": officer(b),
            }
            for b in bs
            if counted(b) and in_range(b.start_date) and unit_ok(b)
            and (not category_id or any(it.asset.category_id == category_id for it in b.items))
        ]
        cols = [("kode", "Kode"), ("peminjam", "Peminjam"), ("unit", "Unit Kerja"), ("barang", "Barang"),
                ("pinjam", "Tgl Pinjam"), ("batas", "Batas Kembali"), ("status", "Status"), ("petugas", "Petugas")]
        summary = f"{len(rows)} transaksi peminjaman"
    elif rtype == "pengembalian":
        rs = session.scalars(select(Return).order_by(Return.returned_at.desc())).all()
        rows = []
        for r in rs:
            b = r.borrowing
            if not in_range(r.returned_at.date()) or not unit_ok(b):
                continue
            det = "; ".join(
                f"{ri.asset.inventory_code}: {ri.asset_status_after or ri.final_condition}"
                + ("" if not ri.completeness or "tidak" not in ri.completeness.lower() else " (tidak lengkap)")
                for ri in r.items
            )
            late = max(0, (r.returned_at.date() - b.due_date).days)
            rows.append({
                "kode": f"KMB-{r.id:04d}",
                "transaksi": b.transaction_number,
                "peminjam": b.borrower.name,
                "kondisi": det,
                "tgl": _d(r.returned_at),
                "terlambat": f"{late} hari" if late else "Tepat waktu",
                "petugas": r.receiver.full_name if r.receiver else "—",
            })
        cols = [("kode", "Kode Kembali"), ("transaksi", "Transaksi"), ("peminjam", "Peminjam"), ("kondisi", "Kondisi Akhir"),
                ("tgl", "Tgl Kembali"), ("terlambat", "Keterlambatan"), ("petugas", "Diterima Oleh")]
        summary = f"{len(rows)} pengembalian"
    elif rtype == "keterlambatan":
        bs = [b for b in session.scalars(select(Borrowing)).all()
              if counted(b) and late_days(b) > 0 and in_range(b.due_date) and unit_ok(b)]
        bs.sort(key=late_days, reverse=True)
        rows = []
        for b in bs:
            events = session.scalars(
                select(NotificationEvent.event_code).where(NotificationEvent.borrowing_id == b.id,
                                                           NotificationEvent.event_code.like("H_PLUS_%"))
            ).all()
            returned = max((r.returned_at for r in b.returns), default=None)
            rows.append({
                "kode": b.transaction_number,
                "peminjam": b.borrower.name,
                "unit": b.borrower.unit.name if b.borrower.unit else "",
                "barang": items_txt(b),
                "batas": _d(b.due_date),
                "kembali": _d(returned) if returned else "Belum kembali",
                "terlambat": f"{late_days(b)} hari",
                "status": display(b),
                "notifikasi": ", ".join(EVENT_TO_UI.get(e, e) for e in events) or "—",
            })
        cols = [("kode", "Kode"), ("peminjam", "Peminjam"), ("unit", "Unit"), ("barang", "Barang"), ("batas", "Batas"),
                ("kembali", "Dikembalikan"), ("terlambat", "Terlambat"), ("status", "Status"), ("notifikasi", "Notifikasi Terkirim")]
        summary = f"{len(rows)} transaksi terlambat"
    else:
        damaged = ("RUSAK", "RUSAK_BERAT", "DALAM_PERBAIKAN", "HILANG")
        assets = session.scalars(select(Asset).where(Asset.deleted_at.is_(None), Asset.is_active.is_(True), Asset.status.in_(damaged)).order_by(Asset.inventory_code)).all()
        rows = []
        for a in assets:
            if not asset_ok(a):
                continue
            h = session.scalars(select(AssetHistory).where(AssetHistory.asset_id == a.id, AssetHistory.new_status == a.status)
                                .order_by(AssetHistory.created_at.desc())).first()
            rows.append({
                "kode": a.inventory_code,
                "nama": a.name,
                "kategori": a.category.name if a.category else "",
                "lokasi": a.location.name if a.location else "",
                "status": a.status,
                "kondisi": CONDITION_LABEL.get(a.condition, a.condition),
                "sejak": _d(h.created_at) if h else "—",
                "keterangan": h.reason if h and h.reason else "",
                "nilai": float(a.acquisition_cost) if a.acquisition_cost is not None else None,
            })
        cols = [("kode", "Kode"), ("nama", "Nama Barang"), ("kategori", "Kategori"), ("lokasi", "Lokasi"), ("status", "Status"),
                ("kondisi", "Kondisi"), ("sejak", "Sejak"), ("keterangan", "Keterangan"), ("nilai", "Nilai Perolehan", True)]
        summary = f"{len(rows)} barang rusak/perbaikan/hilang"

    columns = [{"key": c[0], "label": c[1], "money": len(c) > 2} for c in cols]
    return Report(rtype, REPORT_TYPES[rtype], columns, rows, summary, f)


def _period(f: dict[str, Any]) -> str:
    a, b = _date(f.get("from")), _date(f.get("to"))
    if a and b:
        return f"Periode {_d(a)} s.d. {_d(b)}"
    if a:
        return f"Sejak {_d(a)}"
    if b:
        return f"Sampai {_d(b)}"
    return "Semua periode"


def to_xlsx(rep: Report, settings: dict, user_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = rep.type.capitalize()[:31]
    ncol = len(rep.columns)
    ws.append([f"{rep.title.upper()}"])
    ws.append([f"{settings.get('institution', '')} — {settings.get('unit_sarpras', '')}"])
    ws.append([f"{_period(rep.filters)} · Dicetak {datetime.now():%d/%m/%Y %H:%M} oleh {user_name}"])
    ws.append([])
    for r in range(1, 4):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(1, ncol))
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([c["label"] for c in rep.columns])
    head_row = 5
    thin = Side(style="thin", color="CBD5E1")
    for i in range(1, ncol + 1):
        cell = ws.cell(row=head_row, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in rep.rows:
        ws.append([row.get(c["key"]) if row.get(c["key"]) is not None else "" for c in rep.columns])
    for i, c in enumerate(rep.columns, start=1):
        letter = get_column_letter(i)
        width = max([len(str(c["label"]))] + [len(str(r.get(c["key"]) or "")) for r in rep.rows[:500]])
        ws.column_dimensions[letter].width = min(60, max(10, width + 2))
        if c["money"]:
            for r in range(head_row + 1, head_row + 1 + len(rep.rows)):
                ws.cell(row=r, column=i).number_format = '"Rp" #,##0'
    if rep.rows:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(ncol)}{head_row + len(rep.rows)}"
    ws.freeze_panes = f"A{head_row + 1}"
    ws.append([])
    ws.append([rep.summary])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rep: Report, settings: dict, user_name: str) -> bytes:
    from weasyprint import HTML

    e = html.escape

    def cell(c, r):
        v = r.get(c["key"])
        if c["money"]:
            return _rupiah(v) if v not in (None, "") else "—"
        return e(str(v)) if v not in (None, "") else "—"

    head = "".join(f"<th>{e(c['label'])}</th>" for c in rep.columns)
    body = "".join(
        "<tr>" + "".join(f"<td class='{'num' if c['money'] else ''}'>{cell(c, r)}</td>" for c in rep.columns) + "</tr>"
        for r in rep.rows
    ) or f"<tr><td colspan='{len(rep.columns)}' class='empty'>Tidak ada data</td></tr>"
    doc = f"""<!doctype html><html lang="id"><head><meta charset="utf-8"><style>
    @page {{ size: A4 landscape; margin: 14mm 12mm 16mm; @bottom-right {{ content: "Halaman " counter(page) " dari " counter(pages); font-size: 8pt; color: #64748b; }}
            @bottom-left {{ content: "SIIPB — {e(rep.title)}"; font-size: 8pt; color: #64748b; }} }}
    body {{ font-family: "DejaVu Sans", Arial, sans-serif; font-size: 8.5pt; color: #0f172a; }}
    .kop {{ border-bottom: 2px solid #1e3a8a; padding-bottom: 6px; margin-bottom: 10px; }}
    .kop h1 {{ margin: 0; font-size: 14pt; color: #1e3a8a; letter-spacing: .5px; }}
    .kop div {{ color: #334155; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th {{ background: #1e3a8a; color: #fff; text-align: left; padding: 5px; font-weight: 600; }}
    td {{ border-bottom: 1px solid #e2e8f0; padding: 4px 5px; vertical-align: top; }}
    tr:nth-child(even) td {{ background: #f8fafc; }}
    td.num {{ text-align: right; white-space: nowrap; }}
    td.empty {{ text-align: center; color: #64748b; padding: 16px; }}
    thead {{ display: table-header-group; }} tr {{ page-break-inside: avoid; }}
    .summary {{ margin-top: 8px; font-weight: 600; }}
    .ttd {{ margin-top: 28px; width: 100%; }} .ttd td {{ border: 0; text-align: center; width: 50%; background: none !important; }}
    </style></head><body>
    <div class="kop"><h1>{e(rep.title.upper())}</h1>
    <div>{e(settings.get('institution', ''))} — {e(settings.get('unit_sarpras', ''))}</div>
    <div>{e(_period(rep.filters))} · Dicetak {datetime.now():%d/%m/%Y %H:%M} oleh {e(user_name)}</div></div>
    <table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>
    <div class="summary">{e(rep.summary)}</div>
    <table class="ttd"><tr><td>Mengetahui,<br>Pimpinan<br><br><br><br>(..............................)</td>
    <td>{e(_d(date.today()))}<br>Petugas Sarpras/IT<br><br><br><br>{e(user_name)}</td></tr></table>
    </body></html>"""
    return HTML(string=doc).write_pdf()
