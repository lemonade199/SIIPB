"""Script untuk menggabungkan seluruh modul pembelajaran API SIIPB (01-10) menjadi dokumen Word (.docx) lengkap dan profesional."""
import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT_DIR / "docs" / "api-learning"
OUTPUT_DOCX_ROOT = ROOT_DIR / "PANDUAN-LENGKAP-API-SIIPB.docx"
OUTPUT_DOCX_DOCS = ROOT_DIR / "docs" / "PANDUAN-LENGKAP-API-SIIPB.docx"

# Warna Tema
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy Blue #1E3A8A
COLOR_SECONDARY = RGBColor(14, 116, 144)  # Cyan/Teal #0E7490
COLOR_TEXT = RGBColor(30, 41, 59)         # Slate 800 #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500 #64748B
HEX_PRIMARY = "1E3A8A"
HEX_BG_HEADER = "1E3A8A"
HEX_BG_ALT = "F8FAFC"
HEX_BG_CODE = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BORDER = "0284C7"
HEX_CALLOUT_BG = "F0F9FF"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_inline_text(paragraph, text, base_font="Calibri", base_size=Pt(10), base_color=COLOR_TEXT, italic_default=False):
    """Parse format markdown: **bold**, *italic*, `code`, and plain text."""
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = base_font
        run.font.size = base_size
        run.font.color.rgb = base_color
        if italic_default:
            run.italic = True
            
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_size.pt - 0.5)
            run.font.color.rgb = RGBColor(190, 24, 93)  # Rose/Pink code color
            run.bold = True
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                run.text = m.group(1)
                run.font.color.rgb = RGBColor(2, 132, 199)  # Sky blue link
                run.underline = True
            else:
                run.text = token
        else:
            run.text = token

def add_code_block(doc, code_lines):
    text_content = "\n".join(code_lines)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_BG_CODE)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Border sebelah kiri aksen tebal
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text_content)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    # Jarak setelah blok kode
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_callout_box(doc, quote_lines):
    text_content = " ".join([l.strip() for l in quote_lines if l.strip()])
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="28" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    format_inline_text(p, text_content, base_font="Calibri", base_size=Pt(9.5), base_color=RGBColor(3, 105, 161), italic_default=True)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_markdown_table(doc, header_row, data_rows):
    if not header_row:
        return
    num_cols = len(header_row)
    table = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, color=HEX_BORDER, sz="4")
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(header_row):
        cell = hdr_cells[i]
        set_cell_background(cell, HEX_BG_HEADER)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title.strip())
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Repeat header on new page
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Data rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        # Zebra striping
        bg_col = HEX_BG_ALT if r_idx % 2 == 1 else "FFFFFF"
        for c_idx in range(num_cols):
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            cell = row_cells[c_idx]
            if bg_col != "FFFFFF":
                set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            format_inline_text(p, val.strip(), base_font="Calibri", base_size=Pt(8.0))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def parse_markdown_file_to_doc(doc, md_filepath):
    with open(md_filepath, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    
    in_table = False
    table_header = []
    table_rows = []
    
    in_quote = False
    quote_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\r\n")
        
        # 1. Code Block Handling
        if line.startswith("```"):
            if in_code_block:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # 2. Table Handling
        if "|" in line and (line.strip().startswith("|") or line.strip().endswith("|")):
            parts = [p.strip() for p in line.strip().split("|")]
            # Bersihkan elemen kosong di ujung akibat split (| a | b |)
            if parts and parts[0] == "":
                parts.pop(0)
            if parts and parts[-1] == "":
                parts.pop()
                
            # Cek apakah separator baris (|---|---|)
            if all(re.match(r"^:?-+:?$", p) for p in parts):
                i += 1
                continue
                
            if not in_table:
                in_table = True
                table_header = parts
                table_rows = []
            else:
                table_rows.append(parts)
            i += 1
            continue
        else:
            if in_table:
                add_markdown_table(doc, table_header, table_rows)
                in_table = False
                table_header = []
                table_rows = []

        # 3. Blockquote Handling
        if line.startswith(">"):
            in_quote = True
            quote_lines.append(line.lstrip("> ").strip())
            i += 1
            continue
        else:
            if in_quote:
                add_callout_box(doc, quote_lines)
                in_quote = False
                quote_lines = []

        # Blank line
        if not line.strip():
            i += 1
            continue
            
        # Horizontal Rule
        if re.match(r"^-{3,}$", line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("_________________________________________________________________________________")
            r.font.color.rgb = RGBColor(226, 232, 240)
            r.font.size = Pt(8)
            i += 1
            continue
            
        # 4. Headings
        if line.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[2:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue
            
        if line.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[3:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue
            
        if line.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[4:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(11.5)
            run.bold = True
            run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue
            
        # 5. Bullet Lists
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = line.strip()[2:]
            format_inline_text(p, content, base_size=Pt(9.5))
            i += 1
            continue
            
        # 6. Numbered Lists
        m_num = re.match(r"^(\d+)\.\s+(.*)$", line.strip())
        if m_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            format_inline_text(p, m_num.group(2), base_size=Pt(9.5))
            i += 1
            continue
            
        # 7. Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        format_inline_text(p, line.strip(), base_size=Pt(9.5))
        i += 1

    # Cleanup open blocks at end of file
    if in_code_block:
        add_code_block(doc, code_lines)
    if in_table:
        add_markdown_table(doc, table_header, table_rows)
    if in_quote:
        add_callout_box(doc, quote_lines)

def build_cover_page(doc):
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(70)
    
    # Title Tag Badge
    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = p_tag.add_run("MODUL PEMBELAJARAN & DOKUMENTASI LENGKAP BACKEND REST API")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(10)
    r_tag.bold = True
    r_tag.font.color.rgb = COLOR_SECONDARY
    
    # Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("SIIPB REST API ENGINE")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(28)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Sistem Informasi Inventaris, Peminjaman, dan Pengembalian Barang\nPanduan Komprehensif Arsitektur, 84 Endpoint, Database, Alur Bisnis, dan Latihan PKL")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = COLOR_MUTED
    
    # Horizontal Divider
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("—" * 35)
    r_div.font.color.rgb = RGBColor(203, 213, 225)
    
    # Metadata Box
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(140)
    p_meta.paragraph_format.line_spacing = 1.3
    
    r_meta1 = p_meta.add_run("Target Pembaca: Siswa SMK RPL, Mahasiswa Magang, & Backend Developer\n")
    r_meta1.font.name = "Calibri"
    r_meta1.font.size = Pt(10)
    r_meta1.bold = True
    r_meta1.font.color.rgb = COLOR_PRIMARY
    
    r_meta2 = p_meta.add_run("Stack Teknologi: Python 3.12 · Flask · MariaDB 11 · SQLAlchemy · Alembic · Celery · Next.js · Swagger\nVersi Dokumen: 1.0 (Edisi Lengkap Terverifikasi Source Code)")
    r_meta2.font.name = "Calibri"
    r_meta2.font.size = Pt(9.5)
    r_meta2.font.color.rgb = COLOR_MUTED

def generate_full_docx():
    print("Memulai pembuatan dokumen kompilasi DOCX...")
    doc = Document()
    
    # Konfigurasi Page Setup (A4, margin 0.8 inch agar tabel lebar nyaman)
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4 Width
        section.page_height = Inches(11.69) # A4 Height
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Set default style font
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(9.5)
    doc.styles['Normal'].font.color.rgb = COLOR_TEXT
    
    # 1. Cover Page
    print("Membuat Cover Page...")
    build_cover_page(doc)
    doc.add_page_break()
    
    # Daftar File Markdown
    md_files = [
        "01-dasar-api.md",
        "02-arsitektur-project.md",
        "03-daftar-endpoint.md",
        "04-alur-peminjaman.md",
        "05-alur-pengembalian.md",
        "06-database-dan-api.md",
        "07-auth-dan-role.md",
        "08-error-dan-keamanan.md",
        "09-glosarium.md",
        "10-latihan-presentasi.md",
    ]
    
    for idx, fname in enumerate(md_files):
        fpath = DOCS_DIR / fname
        if not fpath.exists():
            print(f"Warning: File {fname} tidak ditemukan, dilewati.")
            continue
        print(f"Memproses modul {idx+1}/{len(md_files)}: {fname}...")
        parse_markdown_file_to_doc(doc, fpath)
        if idx < len(md_files) - 1:
            doc.add_page_break()
            
    print(f"Menyimpan dokumen ke {OUTPUT_DOCX_ROOT}...")
    doc.save(str(OUTPUT_DOCX_ROOT))
    
    print(f"Menyalin salinan ke {OUTPUT_DOCX_DOCS}...")
    doc.save(str(OUTPUT_DOCX_DOCS))
    
    size_kb = os.path.getsize(str(OUTPUT_DOCX_ROOT)) / 1024
    print(f"SUKSES! Dokumen berhasil dibuat ({size_kb:.1f} KB).")

if __name__ == "__main__":
    generate_full_docx()
