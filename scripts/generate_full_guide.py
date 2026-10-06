"""Script to generate PROJECT-BACKEND-DATABASE-GUIDE.docx
Comprehensive 20-Section Documentation for SIIPB Backend & Database.
"""
import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Colors
COLOR_PRIMARY = RGBColor(27, 54, 93)     # Deep Navy #1B365D
COLOR_SECONDARY = RGBColor(16, 124, 65)  # Emerald #107C41
COLOR_MUTED = RGBColor(100, 116, 139)    # Slate #64748B
COLOR_DARK = RGBColor(30, 41, 59)        # Dark Charcoal #1E293B
HEX_PRIMARY = "1B365D"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_ACCENT_BG = "EFF6FF"
HEX_MUTED_BG = "F1F5F9"

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        is_header = (i == 0)
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
            if col_widths and j < len(col_widths):
                cell.width = Inches(col_widths[j])
            if is_header:
                set_cell_background(cell, HEX_PRIMARY)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, HEX_LIGHT_BG)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = COLOR_DARK

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        for r in h.runs:
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = COLOR_SECONDARY
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = COLOR_DARK
    return h

def add_p(doc, text, bold_prefix="", italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.size = Pt(10)
        r_bold.font.color.rgb = COLOR_DARK
    r_text = p.add_run(text)
    r_text.font.size = Pt(10)
    r_text.italic = italic
    r_text.font.color.rgb = COLOR_DARK
    return p

def add_callout(doc, text, title="CATATAN PENTING"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_ACCENT_BG)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    r_title.font.size = Pt(9.5)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "1E293B")  # Dark code background
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_BORDER}"/><w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/></w:tcBorders>')
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(241, 245, 249)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Framework setup completed.")
